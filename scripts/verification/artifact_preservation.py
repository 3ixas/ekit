#!/usr/bin/env python3
"""Exercise synthetic artifact outcomes with installed libraries and office tools.

Requires python-docx, python-pptx, openpyxl, pypdf, reportlab, soffice and
pdftoppm. Writes only to a newly created explicit output directory. No uploads.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import shutil
import subprocess


def command(args: list[str], log_dir: Path, label: str) -> None:
    if args[0] == "soffice":
        profile = (log_dir.parent / "office-profile").resolve().as_uri()
        args = [args[0], f"-env:UserInstallation={profile}", *args[1:]]
    result = subprocess.run(args, capture_output=True)
    (log_dir / f"{label}.stdout.log").write_bytes(result.stdout)
    (log_dir / f"{label}.stderr.log").write_bytes(result.stderr)
    if result.returncode:
        raise RuntimeError(f"{label} exited {result.returncode}; see {log_dir}")


def run(output: Path) -> dict:
    from docx import Document
    from docx.shared import Inches as DocInches
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.page import PageMargins
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas

    for binary in ("soffice", "pdftoppm"):
        if not shutil.which(binary):
            raise RuntimeError(f"Missing {binary}; rendering/calculation not verified")
    output.mkdir(parents=True, exist_ok=False)
    logs = output / "logs"
    logs.mkdir()
    rendered = output / "rendered"
    rendered.mkdir()
    calculated = output / "calculated"
    calculated.mkdir()

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = DocInches(8.5), DocInches(11)
    document.add_heading("Artifact verification", level=1)
    document.add_paragraph("Synthetic fixture: no personal or employer data.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Measure", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "Expected total", "72"
    document.add_page_break()
    document.add_heading("Second page", level=1)
    document.add_paragraph("Pagination and saved structure are separate checks.")
    docx_path = output / "document.docx"
    document.save(docx_path)
    reopened = Document(docx_path)
    assert reopened.paragraphs[0].style.name == "Heading 1"
    assert reopened.tables[0].cell(1, 1).text == "72"
    command(["soffice", "--headless", "--convert-to", "pdf:writer_pdf_Export",
             "--outdir", str(rendered), str(docx_path)], logs, "document-export")
    doc_pdf = PdfReader(rendered / "document.pdf")
    assert len(doc_pdf.pages) == 2
    assert "Artifact verification" in doc_pdf.pages[0].extract_text()
    assert "Second page" in doc_pdf.pages[1].extract_text()

    source_pdf = output / "source.pdf"
    drawing = canvas.Canvas(str(source_pdf), pagesize=(612, 792))
    drawing.drawString(72, 720, "First source page")
    drawing.showPage()
    drawing.drawString(72, 720, "Second source page - retained")
    drawing.showPage()
    drawing.save()
    reader, writer = PdfReader(source_pdf), PdfWriter()
    writer.add_page(reader.pages[1].rotate(90))
    rotated = output / "rotated.pdf"
    writer.write(rotated)
    transformed = PdfReader(rotated)
    assert len(transformed.pages) == 1
    page = transformed.pages[0]
    assert page.rotation == 90
    assert list(page.mediabox) == [0, 0, 612, 792]
    assert "Second source page - retained" in page.extract_text()
    assert "First source page" not in page.extract_text()

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for title, body in [("Make the outcome testable", "Two editable slides, two inspected renders."),
                        ("Known totals provide an independent check", "24 + 48 = 72")]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        heading = slide.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(11.8), Inches(1.4))
        heading.text_frame.text = title
        heading.text_frame.paragraphs[0].font.size = Pt(32)
        content = slide.shapes.add_textbox(Inches(0.7), Inches(2.5), Inches(11.8), Inches(2))
        content.text_frame.text = body
        content.text_frame.paragraphs[0].font.size = Pt(25)
    pptx_path = output / "presentation.pptx"
    deck.save(pptx_path)
    reopened_deck = Presentation(pptx_path)
    assert len(reopened_deck.slides) == 2
    for slide in reopened_deck.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width
            assert shape.top + shape.height <= deck.slide_height
    command(["soffice", "--headless", "--convert-to", "pdf:impress_pdf_Export",
             "--outdir", str(rendered), str(pptx_path)], logs, "slide-export")
    slides_pdf = PdfReader(rendered / "presentation.pdf")
    assert len(slides_pdf.pages) == 2
    assert "Make the outcome testable" in slides_pdf.pages[0].extract_text()
    assert "72" in slides_pdf.pages[1].extract_text()

    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    for row in [("Item", "Quantity", "Price", "Subtotal"),
                ("A", 2, 12, "=B2*C2"), ("B", 6, 8, "=B3*C3"),
                ("Total", None, None, "=SUM(D2:D3)")]:
        sheet.append(row)
    for name in "ABCD":
        sheet.column_dimensions[name].width = 20
    sheet.freeze_panes = "A2"
    sheet.print_area = "A1:D4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 1
    sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.3, bottom=0.3)
    for row in sheet.iter_rows(min_row=2, max_row=4, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = "0.00"
    workbook_path = output / "budget.xlsx"
    book.save(workbook_path)
    # Negative case: writing a formula has not evaluated it.
    assert load_workbook(workbook_path, data_only=True)["Budget"]["D4"].value is None
    command(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
             "--outdir", str(calculated), str(workbook_path)], logs, "workbook-recalculate")
    result_path = calculated / "budget.xlsx"
    values = load_workbook(result_path, data_only=True)["Budget"]
    formulas = load_workbook(result_path, data_only=False)["Budget"]
    assert [values[cell].value for cell in ("D2", "D3", "D4")] == [24, 48, 72]
    assert formulas["D2"].value == "=B2*C2"
    assert formulas["D4"].value == "=SUM(D2:D3)"
    assert formulas["D4"].number_format == "0.00"
    assert not any(cell.data_type == "e" for row in values for cell in row)
    command(["soffice", "--headless", "--convert-to", "pdf:calc_pdf_Export",
             "--outdir", str(rendered), str(result_path)], logs, "workbook-export")
    budget_pdf = PdfReader(rendered / "budget.pdf")
    assert len(budget_pdf.pages) == 1
    assert "72" in budget_pdf.pages[0].extract_text()

    for source, name in [(rendered / "document.pdf", "document"),
                         (rotated, "rotated"),
                         (rendered / "presentation.pdf", "slides"),
                         (rendered / "budget.pdf", "budget")]:
        command(["pdftoppm", "-scale-to", "1100", "-png", str(source), str(rendered / name)],
                logs, f"render-{name}")
    result = {"structural_and_numeric_checks": "pass",
              "document_pages": 2, "rotated_pages": 1, "slides": 2,
              "workbook_total": 72, "formula_only_write_is_not_recalculation": True,
              "images": sorted(str(p.relative_to(output)) for p in rendered.glob("*.png")),
              "visual_inspection": "pending human/agent inspection of every listed image",
              "scope": "Synthetic static fixtures only; no revision/form/OCR/media/accessibility conformance claim."}
    (output / "results.json").write_text(json.dumps(result, indent=2) + "\n")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(run(args.output_dir.resolve()), indent=2))


if __name__ == "__main__":
    main()
